from docx import Document
from pptx.util import Inches
from pptx.util import Pt
from pptx.dml.color import RGBColor
import copy

from handlers.handle_tags import handle_tags

def add_data_to_table(table):
    data = (
        (1, 'gaya 1'),
        (2, 'gaya 2'),
        (3, 'gaya 3'),
        (1, 'gaya 4'),
        (2, 'gaya 5'),
        (3, 'gaya 6'),
        (1, 'gaya 7'),
        (2, 'gaya 8'),
        (3, 'gaya 9'),
        (1, 'gaya 10'),
        (1, 'gaya 11'),
        (2, 'gaya 12'),
        (3, 'gaya 13'),
        (1, 'gaya 14'),
        (2, 'gaya 15'),
        (3, 'gaya 16'),
        (1, 'gaya 17'),
        (2, 'gaya 18'),
        (3, 'gaya 19'),
        (1, 'gaya 20')
    )

    print(table.rows[0].cells)

    row = table.rows[0].cells
    row[0].text = 'Id'
    row[1].text = 'Name'
    row[2].text = 'aaa'
    row[3].text = 'bbbb'
    row[4].text = 'cccc'

    for id, name in data:
        row = table.add_row().cells
        row[0].text = str(id)
        row[1].text = "name"
        row[2].text = "name 1"
        row[3].text = "nam  2"
        row[4].text = "name 3"

    print("table style===========>",table.style)

def update_table(replacements, document):  
    print("template table style ===========>",document.tables[0].style)
    #add data to existing table
    add_data_to_table(document.tables[0])

    new_tbl = copy.deepcopy(document.tables[0])
    print("copied table============>",new_tbl)

    #table = document.add_table(rows=1, cols=6, style=)
    paragraph = doc.add_paragraph(' ')
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(5)

    #create new table using existing styles
    table = document.add_table(rows=1, cols=5)
    table.style = document.tables[0].style
    add_data_to_table(table)
    paragraph = doc.add_paragraph(' ')
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(5)


    table_2 = document.add_table(rows=1, cols=5)
    table.style = 'Table Grid'
    add_data_to_table(table_2)

if __name__ == '__main__':

    doc = Document('input.docx')
    dataDump = {
        "text_replaces":  {
            'var1': "Example Name",
            'var2': 'External asset workflow',
            'var3': "302929393",
            'var4': 'LGIM/Client user',
            'var5': "+972-5056000000",
            'var6': "example@example.com",
            'var7': 'assets',
            'var8': 'Information button text',
            'var10': "03 Jan, 2022",
            'var11': "10,000",
            'var12': "3,000",
            'var13':"7,000",
            'spart': "sample text",
            'scustomer': "sample customer",
            'sname': "test",
            'suser': "gayashan",
            'execSum': "Contrary to popular belief, Lorem Ipsum is not simply random text. It has roots in a piece of classical Latin literature  from 45 BC, making it over 2000 years old. Richard McClintock, a Latin professor at Hampden-Sydney College in Virginia, looked up one of the more obscure Latin words, consectetur, from a Lorem Ipsum passage, and going through the cites of the word ",
            'content1': "Content sample 1",
            'sect': "Section 1",
            'content_summary': 'Content',
            'sec' : 'Section two',
            'sec_2_1': 'Section two one',
            'sec_2_1_1': 'Section two one one',
            'summary_2_1_2': 'Summary two one two',
            'summary_2_2': 'Summary two two',

            },
        "table_replaces": {
            "cashFlows":{
                "headers": ["cashflow year","cashflow fixed","cashflow real"],
                "rows": [ 
                    {
                        "cashflow_year": "2020","cashflow_fixed":"1","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"2","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"3","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"4","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"5","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"6","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"7","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"8","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"9","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"10","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"11","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"12","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"13","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"14","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"15","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"16","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"17","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"18","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"19","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"20","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"21","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"22","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"23","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"24","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"25","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"26","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"27","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"28","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"29","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"30","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"31","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"32","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"33","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"34","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"35","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"36","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"37","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"38","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"39","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"40","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"41","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"42","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"43","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"44","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"45","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"51","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"52","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"53","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"54","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"55","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"56","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"57","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"58","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"59","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"510","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"151","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"512","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"153","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"154","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"155","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"61","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"62","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"63","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"64","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"65","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"66","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"67","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"68","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"69","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"170","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"171","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"172","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"173","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"174","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"175","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"17","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"82","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"83","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"84","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"85","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"86","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"87","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"88","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"89","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"190","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"191","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"192","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"193","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"194","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"195","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"17","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"82","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"83","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"84","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"85","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"86","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"87","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"88","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"89","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"190","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"191","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"192","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"193","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"194","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"195","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"1","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"2","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"3","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"4","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"5","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"6","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"7","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"8","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"9","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"10","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"11","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"12","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"13","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"14","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"15","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"16","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"17","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"18","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"19","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"20","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"21","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"22","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"23","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"24","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"25","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"26","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"27","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"28","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"29","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"30","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"31","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"32","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"33","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"34","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"35","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"36","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"37","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"38","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"39","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"40","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"41","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"42","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"43","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"44","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"45","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"51","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"5112","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"5113","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"5114","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"5115","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"5116","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"5117","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"58","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"5911","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"51110","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"15111","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"5112","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"1513","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"1514","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"1515","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"611","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"612","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"6113","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"614","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"615","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"616","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"617","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"618","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"619","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"1170","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"1171","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"1172","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"1173","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"1714","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"1715","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"1117","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"812","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"813","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"814","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"815","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"816","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"817","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"818","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"819","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"1910","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"1911","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"1912","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"1913","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"1914","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"1915","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"117","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"812","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"813","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"814","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"815","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"816","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"187","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"818","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"189","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"1190","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"1911","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"1192","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"1913","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"1941","cashflow_real": "123123"
                    },
                    {
                        "cashflow_year": "2020","cashflow_fixed":"1915","cashflow_real": "123123"
                    }
                 ]
            }
        } ,
        "table_1_data": {
            "data":[{
                "Header 1": "data 1","Header 2": "data 2","Header 3": "data 3","Header 4": "data 4","Header 5": "data 5"#,"Header 6": "data 6","Header 7": "data 7","Header 8": "data 8"
            },
            {
                "Header 1": "data 11","Header 2": "data 12","Header 3": "data 13","Header 4": "data 14","Header 5": "data 15"#,"Header 6": "data 16","Header 7": "data 17","Header 8": "data 18"
            },
            {
                "Header 1": "data 21","Header 2": "data 22","Header 3": "data 23","Header 4": "data 24","Header 5": "data 25"#,"Header 6": "data 26","Header 7": "data 27","Header 8": "data 28"
            },
            {
                "Header 1": "data 31","Header 2": "data 32","Header 3": "data 33","Header 4": "data 34","Header 5": "data 35"#,"Header 6": "data 36","Header 7": "data 37","Header 8": "data 38"
            },
            {
                "Header 1": "data 41","Header 2": "data 42","Header 3": "data 43","Header 4": "data 44","Header 5": "data 45"#,"Header 6": "data 46","Header 7": "data 47","Header 8": "data 48"
            },
            {
                "Header 1": "data 51","Header 2": "data 52","Header 3": "data 53","Header 4": "data 54","Header 5": "data 55"#,"Header 6": "data 56","Header 7": "data 57","Header 8": "data 58"
            },
            {
                "Header 1": "data 61","Header 2": "data 62","Header 3": "data 63","Header 4": "data 64","Header 5": "data 65"#,"Header 6": "data 66","Header 7": "data 67","Header 8": "data 68"
            },
            {
                "Header 1": "data 71","Header 2": "data 72","Header 3": "data 73","Header 4": "data 74","Header 5": "data 75"#,"Header 6": "data 76","Header 7": "data 77","Header 8": "data 78"
            },
            {
                "Header 1": "data 81","Header 2": "data 82","Header 3": "data 83","Header 4": "data 84","Header 5": "data 85"#,"Header 6": "data 86","Header 7": "data 87","Header 8": "data 88"
            },
            {
                "Header 1": "data 91","Header 2": "data 92","Header 3": "data 93","Header 4": "data 94","Header 5": "data 95"#,"Header 6": "data 96","Header 7": "data 97","Header 8": "data 98"
            }
            ]
        }
             
    }

    handle_tags(doc,dataDump)

    doc.save('output.docx')